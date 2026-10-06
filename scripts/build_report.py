"""Build the paper-style manuscript from frozen, executed evidence."""
from __future__ import annotations
import json
import math
import re
import statistics
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'report'; OUT.mkdir(exist_ok=True)


def load(name):
    return json.loads((ROOT/name).read_text(encoding='utf-8'))


def main():
    sim=load('results/simulation_summary.json');hist=load('results/historical_summary.json')
    parallel=load('results/parallel_benchmark.json');union=load('docs/competency_union.json');refs=load('docs/references.json')
    validation=load('results/validation.json') if (ROOT/'results/validation.json').exists() else {}
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.78);section.bottom_margin=Inches(.72)
    section.left_margin=section.right_margin=Inches(.85)
    for name in ['Normal','Body Text','Caption','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ['Normal','Body Text','Caption'] else '微软雅黑')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.22;normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',24),('Subtitle',12),('Heading 1',15),('Heading 2',12),('Heading 3',11),('Caption',9)]:
        doc.styles[name].font.size=Pt(size)
    for name in ['Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.bold=True;st.paragraph_format.space_before=Pt(15);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
    for st in doc.styles:
        for border in list(st._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Subtitle'].font.italic=False
    doc.core_properties.title='有限数据与模型不确定性下的期权对冲研究'
    doc.core_properties.subject='金融数学课程延伸计算研究与十项目能力并集证据'
    doc.core_properties.author='Project initiated by Youming Yang; AI-assisted development and drafting'
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    def p(text,style=None): return doc.add_paragraph(text,style)
    def h(text,level=1): return doc.add_heading(text,level)
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        if widths:
            for c,w in zip(t.columns,widths): c.width=Inches(w)
        for cell,text in zip(t.rows[0].cells,headers):
            cell.text=str(text)
            for run in cell.paragraphs[0].runs: run.bold=True
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEFF1');cell._tc.get_or_add_tcPr().append(shade)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for cell,value in zip(cells,row): cell.text=str(value)
        for row in t.rows:
            no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
            if widths:
                for cell,w in zip(row.cells,widths): cell.width=Inches(w)
            for cell in row.cells:
                tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D0D4D8');borders.append(e)
                tcpr.append(borders)
                for par in cell.paragraphs:
                    par.paragraph_format.space_after=Pt(4);par.paragraph_format.space_before=Pt(3);par.paragraph_format.line_spacing=1.12
                    for run in par.runs:run.font.size=Pt(9)
        p('')
        return t
    def fig(name,caption):
        par=p('');par.paragraph_format.keep_with_next=True
        pic=par.add_run().add_picture(str(ROOT/'figures'/name),width=Inches(6.65))
        pic._inline.docPr.set('descr',caption)
        c=p(caption,'Caption');c.paragraph_format.space_after=Pt(10)
    def equation(text):
        par=p('');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        eq=OxmlElement('m:oMathPara');mathnode=OxmlElement('m:oMath')
        def mr(value):
            run=OxmlElement('m:r');txt=OxmlElement('m:t');txt.text=value;run.append(txt);return run
        cursor=0
        for match in re.finditer(r'([A-Za-zΑ-Ωα-ω])([_^])(?:\{([^}]+)\}|([A-Za-z0-9]))',text):
            if match.start()>cursor:mathnode.append(mr(text[cursor:match.start()]))
            sub=OxmlElement('m:sSub' if match.group(2)=='_' else 'm:sSup');base=OxmlElement('m:e');base.append(mr(match.group(1)));index=OxmlElement('m:sub' if match.group(2)=='_' else 'm:sup');index.append(mr(match.group(3) or match.group(4)));sub.extend([base,index]);mathnode.append(sub);cursor=match.end()
        if cursor<len(text):mathnode.append(mr(text[cursor:]))
        eq.append(mathnode);par._p.append(eq)
        par.paragraph_format.space_before=Pt(5);par.paragraph_format.space_after=Pt(8)
    def link(par,label,url):
        node=OxmlElement('w:hyperlink');node.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245B70');rp.append(color);run.append(rp);txt=OxmlElement('w:t');txt.text=label;run.append(txt);node.append(run);par._p.append(node)
    def metric(currency,strategy,fee=5):
        return next(x for x in hist['currencies'][currency]['hedging_metrics'] if x['split']=='test' and x['strategy']==strategy and x['cost_bps']==fee)
    def simrow(scenario,policy,fee=5):
        return next(x for x in sim['rows'] if x['scenario']==scenario and x['policy']==policy and x['fee_bps']==fee and x['split']=='test')
    primary=next(x for x in sim['paired_comparisons'] if x['primary'])
    p('有限数据与模型不确定性下的期权对冲研究','Title')
    p('统计估计 数值计算与成本约束决策','Subtitle')
    p('FinMath Risk Lab   课程延伸计算研究报告   2026年10月6日')
    p('项目发起人  杨又铭   版本 1.0   配套成果  Python研究库 本地交互工作台 可复现实验与数据证据')
    h('摘要')
    p('金融数学课程给出了理想市场中的定价和复制框架，但有限数据、离散交易以及模型错误会共同影响实际计算结果。本研究建立从公开数据、方差预测、期权定价到自融资对冲账本的统一实验链，考察更精确的计算、更复杂的估计与更频繁的调仓，何时能够改善终端风险，何时不能。研究同时采用可控路径模拟与欧洲央行历史参考汇率，避免只在单一模型内部验证模型本身。')
    p(f'模拟覆盖七种场景、三档费用和六种策略，每场景使用1024条验证路径与8192条独立测试路径。模型匹配且费用为5 bps时，验证选择的每日策略相对每五步策略的平均平方终端损失差为{primary["difference"]:.4f}，95%配对区间为[{primary["ci_low"]:.4f}, {primary["ci_high"]:.4f}]。这一结论不跨越至全部模型：波动率切换与跳跃场景仍保留显著复制误差，且共同初始权利金的定价失配也是损益来源。')
    p('历史研究采用每欧元对应美元与日元的两条参考汇率，将2005至2016年用于训练、2017至2020年用于验证、2021至2025年用于测试。Ridge在日元系列的方差代理损失上优于滚动基线，但其对冲平方损失差的区间仍包含零。美元系列的阈值策略降低了费用，同时增加净损益RMSE。由此，预测改善、成本下降和最终风险改善应分别检验。')
    p('成果包含解析与数值定价、因果时间特征、冻结模型与验证选择、配对推断、移动块重采样、成本约束、审计账本、SQLite证据库、交互界面和可复现报告。研究贡献是可核验的方法整合与条件性发现，不主张新定价理论、市场最优策略或实盘盈利。开发和报告整理有实质性AI协助，个人贡献与后续学习应另行真实记录。')
    p('关键词  金融数学  波动率预测  离散对冲  交易成本  模型风险  可复现研究')
    h('Abstract')
    p('This course-extension computational study examines the connection between finite-data variance forecasting and cost-aware option hedging. A unified research pipeline links data provenance, analytic and numerical pricing, causal estimation, self-financing cash accounting and risk evaluation. Seven simulated mechanisms are evaluated with separate validation and test paths. A second study uses European Central Bank reference exchange rates with chronological training, validation and test periods. Model and policy selection is kept separate from final evaluation.')
    p('The results support conditional rather than universal improvements. More frequent rebalancing reduces replication loss in a matched low-cost setting. Under misspecification, substantial residual risk persists. Better out-of-sample variance-proxy forecasts for one currency do not establish a statistically clear improvement in downstream hedging loss. A transaction-cost reduction can also increase net replication RMSE. The package includes executed experiments, inspectable fitted parameters, numerical and accounting tests, moving-block uncertainty analysis, a research dashboard and reproducible evidence. Historical experiments use hypothetical options on informational reference prices and do not establish executable trading performance.')
    h('1 研究动机与问题')
    p('本项目延伸自Introduction to Financial Mathematics课程。所提供资料列示课程时间为2024年10月11日至2025年1月5日、54小时、成绩84.40，内容包括货币时间价值、债券利率、远期与期货、无套利、二叉树、动态规划、随机游走、Black–Scholes、波动率、均值方差与CAPM。课程为本研究提供概念起点；本研究没有把所有课程章节扩成互不相关的功能，而是围绕期权复制建立一条可追溯的实证链。证书本身不构成授课教师所在大学的正式学分或研究任职证明。')
    p('计算一个期权价格并不足以回答风险决策问题。解析公式依赖波动率等不可直接观察的输入；用有限历史样本估计输入会产生不确定性；离散调仓又将估计误差与路径变化、费用共同传递到终端损益。如果只展示预测精度或界面功能，便无法知道模型是否改善了下游决策。因此本研究把账本作为所有方法的共同评价终点。')
    table(['研究问题','可观察证据','不能由该证据推出的结论'],[
        ['Q1 数值近似是否正确','解析基准、树收敛、MC标准误、恒等式测试','真实市场价格一定服从模型'],
        ['Q2 有限数据与失配如何影响复制','独立校准样本、制度切换、跳跃与完整损益','各现实风险来源已被完全识别'],
        ['Q3 预测改善是否传到对冲','同日期方差代理损失与共同保费下终端损失','更复杂模型必然经济价值更高'],
        ['Q4 成本预算如何改变策略','验证集有限候选选择及测试风险成本','获得所有策略空间的全局最优']
    ],[1.65,2.45,2.55])
    h('2 相关方法与研究定位')
    p('Black与Scholes的解析定价提供可验证基准[R1]；Cox、Ross与Rubinstein的重组二叉树将定价转化为向后递推[R2]。本项目实现这些已有方法，并用独立金融恒等式和误差趋势检查实现。美式树作为数值能力扩展保留在库中，核心对冲研究统一使用欧式看涨期权，避免混入提前行权政策差异。')
    p('交易成本使连续复制的理想逻辑发生改变。Leland的工作是本研究讨论费用与调仓频率的文献背景[R3]；本项目直接计算离散现金流，没有实现或声称验证其渐近修正定价定理。针对预测评价，Patton强调波动率代理和损失函数选择的重要性[R4]。这里用次观测平方对数收益作为噪声代理，同时报告QLIKE与MSE，不把该代理直接称为真实潜在方差。')
    p('历史序列上的损失比较采用Künsch相关框架所启发的移动块bootstrap[R5]，保留块内短期依赖并展示块长敏感性。Ridge使用公开文档所规定的正则化目标[R6]；时间评估遵守先训练、再验证、最后测试的顺序[R7]。这些方法的组合形成一项课程延伸研究，而非新的学习算法或未经证明的理论突破。')
    h('3 数学设定与账本约定')
    h('3.1 定价与敏感度',2)
    p('记S为现价，K为执行价，τ为剩余期限，r为连续复利率，q为连续收益率，σ为年化波动率，Φ为标准正态分布函数。常参数欧式看涨期权的解析价格与delta为：')
    equation('C = S exp(−qτ) Φ(d₁) − K exp(−rτ) Φ(d₂)')
    equation('d₁ = [ln(S/K) + (r − q + σ²/2)τ] / (σ√τ)     d₂ = d₁ − σ√τ')
    equation('Δ = exp(−qτ) Φ(d₁)')
    p('定价API同时支持看跌与q，并输出gamma、vega、theta、rho。Vega按波动率变化1.0计，rho按利率变化1.0计，theta按一年时间衰减计；界面明确单位。到期、零波动率及不可微折点单独处理。隐含波动率先核验无套利价格界，再进行数值求根，边界情况不伪造有限解。对冲研究则明确限制q=0。')
    h('3.2 二叉树与Monte Carlo',2)
    p('CRR取u=exp(σ√Δt)、d=1/u，以风险中性概率p=(exp((r−q)Δt)−d)/(u−d)向后递推。若网格导致p不在[0,1]内，程序拒绝输入而非截断概率。美式期权在每个节点比较继续持有价值与立即行权价值。该结构直接对应课程中的多步二叉树与动态规划。')
    equation('V = exp(−rT) E^Q[max(S_T − K, 0)]')
    p('Monte Carlo在风险中性终端分布下估价，标准误来自独立估计单位。反变量方法令正态样本Z与−Z成对，以两者折现收益的均值作为一个独立单位。因此200000条终端路径只有100000个独立配对单位。不能以相关的单条路径直接计算独立样本标准误。数值误差不会在每一个随机样本量上严格单调下降。')
    h('3.3 自融资离散对冲',2)
    p('研究从卖方视角出发：初始收到共同权利金，购入delta数量标的，余额存入或借入现金账户。每一步先让上一步现金计息，再按当期可用信息决定持仓，再经历下一步价格变化。交易额绝对值乘费用率得到手续费。入场、调仓与到期平仓全部计费，到期扣除期权赔付。借贷使用同一利率，未引入外部注资、破产停止或保证金。')
    equation('Bₜ = Bₜ₋₁ exp(rΔt) − (Δₜ − Δₜ₋₁)Sₜ − κ |Δₜ − Δₜ₋₁| Sₜ')
    equation('Π_T = B_{N−1} exp(rΔt) + Δ_{N−1}S_T − κ |Δ_{N−1}|S_T − max(S_T − K, 0)')
    p('第二式描述第N步到期时的终端P&L，T=NΔt；代码账本用明确的step与event区分调仓和到期事件。费用总量按融资利率积累到终点。对于不随现金变化的同一交易规则，免费P&L与收费P&L之差应恰等于费用终值，这是比单看回测图更强的核验条件。逐笔账本还记录目标delta、成交量、利息、余额、费用与最终赔付。')
    h('3.4 风险目标与统计单位',2)
    equation('L = −Π     RMSE = √mean(Π²)     QLIKE = mean[ln(h) + y/h]')
    p('RMSE以零终端P&L为目标，包含平均偏差和离散风险，不等于去均值标准差。VaR95采用线性经验分位数；ES95对最坏5%的经验质量取平均，在尾部边界使用部分权重，避免样本量不整除时扩大尾部比例。QLIKE省略仅依赖真实目标的项，数值可以为负，比较必须使用相同观测日期。')
    p('模拟的独立单位是完整测试路径，同一路径上策略结果配对。历史的独立性不作相同假设：预测误差按观测日块重采样，非重叠对冲episode按有序episode块重采样。非重叠不等于相互独立；两个汇率也不合并为独立重复实验。')
    h('4 模拟实验协议')
    p('所有场景统一S₀=K=100、T=0.5年、r=0.03、126个调仓步长，基础波动率0.20，物理漂移0.03。物理漂移与r数值相同是实验设定，不意味着风险中性概率与物理概率概念相同。每场景1024条验证路径与8192条测试路径来自不同随机流；所有候选策略在同一组路径上比较。')
    table(['场景','路径生成与估计','目的'],[
        ['S1','GBM σ=0.20 对冲σ=0.20','模型匹配基准'],['S2','GBM σ=0.20 对冲σ=0.15','低估波动率'],['S3','GBM σ=0.20 对冲σ=0.30','高估波动率'],['S4','20个独立正态历史收益估计σ','短样本估计不确定性'],['S5','120个独立正态历史收益估计σ','较长样本对照'],['S6','中途σ由0.20切换0.40 对冲仍0.20','制度变化失配'],['S7','补偿跳跃 λ=1 均值−0.10 标准差0.15','连续模型遗漏跳跃']
    ],[.55,3.8,2.3])
    p('S4与S5逐路径使用独立于未来路径的校准样本，以样本标准差和252观测年约定年化。两个场景使用不同路径和种子，因此样本长度之间的比较不是配对估计。S6的切换时间由实验者预设，策略不提前知道新增波动率。跳跃过程包含补偿项以维持所规定的期望增长参数。')
    p(f'所有策略和场景共用Black–Scholes基础σ=0.20所给初始权利金{sim["premium"]:.6f}。对跳跃或切换场景，这不一定是公平价格；因此其终端损益同时包含初始定价失配与后续复制误差，不可全归因于调仓规则。采用共同融资是为了避免通过不同初始价格暗中改善某一策略表现。')
    table(['规则类别','候选','验证选择与预算'],[
        ['固定间隔','daily 每步 every5 每5步 every21 每21步','在平均费用终值≤0.50的候选中选择验证RMSE最小者'],
        ['Delta阈值','band02 0.02 band05 0.05 band10 0.10','阈值指目标与当前持仓delta之差 入场和退出固定执行'],
        ['费用情景','0 5 20 bps','无可行候选则选最低成本并显式记录不满足预算']
    ],[1.0,2.4,3.25])
    p('协议哈希在结果文件中保留；协议在当次计算之前写入，不声称第三方预注册。主比较事先限定为S1、5 bps下验证所选策略相对every5的平均平方损失差，使用95%配对t区间。其余20组区间作为探索性结果，没有多重比较后的整体覆盖保证。预算约束是验证平均量，不保证每条路径或测试平均值满足。')
    h('5 真实数据与预测方法')
    h('5.1 数据来源与日期范围',2)
    p('数据来自欧洲央行公开欧元参考汇率[D1]。原始ZIP按来源原样归档，处理CSV选取USD和JPY列并固定分析截止日2025年12月31日，共6913个发布日期，始于1999年1月4日。计价方向均为一欧元对应的外币数量。没有重复日期、所选报价缺失或非正报价；最大自然日间隔5天，超过3天的间隔57次。未插值、前填或对极端收益进行截尾。')
    p('研究将相邻发布观测作为一步，用252步表示一年。周末与节假日的不等自然时间间隔仍是近似来源。现今下载可能包含修订后的历史值，并非逐日封存的数据版本；严格的代码时间顺序无法消除未知来源修订。ECB数据再利用条件、来源标注和加工说明随包保留[D2]，本研究结果不代表ECB分析或认可。')
    table(['阶段','目标日期范围','每货币预测行数','用途'],[
        ['预热','1999至2004年','不作预测性能样本','初始化滚动与EWMA状态'],['训练','2005至2016年','3073','拟合标准化 回归与校正 常量基线'],['验证','2017至2020年','1022','选择超参数 预测族与调仓阈值'],['测试','2021至2025年','1281','冻结选择后的最终评价']
    ],[.7,1.45,1.2,3.3])
    p('特征行时点t预测下一观测的平方对数收益，按目标日期分组。例如2016年最后一个决策预测2017年第一个收益，该行属于验证而不能进入训练。训练CSV中的拟合值标记为样本内结果，不用于泛化证据。')
    h('5.2 四类方差预测',2)
    p('训练常量以训练期平方收益均值预测未来；Rolling63使用截至当前的63步平方收益均值；EWMA按λ加权更新，候选λ为0.94、0.97、0.99。Ridge以次步平方收益对数为训练响应，七个特征是截至当前的1、5、21、63、126步平方收益均值对数、5步收益和、21步负收益比例。标准化参数完全由训练期确定。')
    p('Ridge的L2惩罚α从0.1、1、10、100中按验证QLIKE选择。指数变换后的预测乘以训练期mean(y/exp(fitted_log))，以缓解对数变换后的尺度偏差。这个训练期均值比校正是建模近似，不保证每个条件状态下无偏。方差上下界预设为10⁻¹⁰与0.04；正式结果显示各模型输出均未触及裁剪。参数、系数、截距、标准化与校正量以JSON保存，便于直接审查。')
    p('训练后不再重新拟合系数。测试中滚动窗口与EWMA状态可吸收已经发生的观测，这是因果更新，不能混同为使用未来信息。验证阶段两种货币均选择EWMA λ=0.97为运行预测族；Ridge内部最佳α分别为USD的100和JPY的0.1。即使测试期Ridge点估计更好，也不追改验证选择。')
    h('5.3 历史对冲与区间',2)
    p('每个货币和分组构造非重叠的21收益episode，每段需22个报价，相邻段仅共享边界价格。验证各48段、测试各60段；测试从2021年1月4日起，最后完整段截至2025年12月1日。余下20个测试收益不足一段而不用于对冲，但预测评价保留至年底。每段起始现价和执行价归一为100，r=q=0，T=21/252，初始权利金统一使用起始Rolling63波动率。')
    p('次步方差预测乘252再开根号，作为剩余期限delta的平坦波动率输入。这是简化决策规则，不是完整期限结构预测，也未按真实外汇双利率估值。比较冻结起始波动率、动态Rolling63、EWMA、Ridge每日策略，以及验证所选预测族的每日、每五步和阈值策略。selected_daily与ewma_daily相同，不能计成独立模型。')
    p('费用为0、5、10 bps的假设情景。阈值在验证集5 bps时从0.025、0.05、0.10、0.15、0.20选择，以平均成本≤0.12下最小净RMSE为目标；选定后用于全部测试费用情景。USD选0.10、JPY选0.05。预测损失差采用2000次移动块重采样，块长5、21、63；episode差采用块长1、3、6。区间均为条件于已选模型的逐项近似，不包括选择不确定性。')
    h('6 数值核验与模拟结果')
    h('6.1 解析基准与数值误差',2)
    p(f'在S=K=100、T=1、r=0.05、σ=0.20条件下，解析看涨价格为{sim["numerical"]["analytic_price"]:.8f}。CRR使用16至1024步的偶数网格；Monte Carlo使用2000至200000条终端路径。图1分别展示树误差和Monte Carlo标准误。偶数子序列的收敛不能据此表述成所有网格的严格单调定理。单次计时仅作描述，不用于宣称算法普遍更快。')
    fig('01_numerical.png','图1 解析基准下的数值近似与Monte Carlo不确定性。反变量标准误基于配对均值，横轴为总终端路径数。来源 numerical_benchmarks.csv。')
    h('6.2 七种机制下的风险',2)
    rows=[]
    for selection in sim['selections']:
        if selection['fee_bps']==5:
            r=simrow(selection['scenario'],selection['selected_policy']);rows.append([selection['scenario'],selection['selected_policy'],f'{r["rmse"]:.4f}',f'{r["es95"]:.4f}',f'{r["mean_cost"]:.4f}'])
    table(['场景','验证所选策略','测试RMSE','测试ES95','平均成本'],rows,[.6,1.8,1.4,1.4,1.45])
    fig('02_scenario_risk.png','图2 全部126个测试条件的净损益RMSE。七种场景、六种策略、三档费用均展示，避免只保留表现最佳的条件。每条件8192条路径，同一场景的条件共享路径。')
    p('有限校准样本增加了对冲输入的不确定性，较长校准样本在此实验中减轻该效应，但不能把两种独立场景的差当作配对显著性结果。制度切换和跳跃导致的误差远大于模型匹配基准，说明把数值求解器算得更精确，并不会自动消除生成机制失配。共同初始权利金也使场景间损益包含定价偏差。')
    h('6.3 成本约束与主要比较',2)
    daily=simrow('S1','daily');every5=simrow('S1','every5')
    p(f'S1、5 bps的每日策略测试RMSE为{daily["rmse"]:.6f}，every5为{every5["rmse"]:.6f}。主要配对平方损失差为{primary["difference"]:.6f}，标准误{primary["standard_error"]:.6f}，95%区间[{primary["ci_low"]:.6f}, {primary["ci_high"]:.6f}]。这是对指定路径机制、费用与候选集的证据，不能扩展为市场中每日调仓恒优。')
    fig('03_cost_risk.png','图3 模型匹配场景的费用与风险。竖虚线是0.50平均成本预算；预算用于验证选择，图展示测试结果。同一费率下各点使用同组测试路径。')
    p('高费用下预算可迫使选择较低频的策略，从而增加相对不受相同可行性限制的高频基线的误差。此类比较显示成本约束的代价，不能解释为所选策略在同一可行集合内一定失败。文件逐一保留验证是否可行及测试是否满足预算，未把预算外的结果隐藏。')
    fig('04_paired_comparisons.png','图4 5 bps下所选策略相对every5的平均平方损失差及95%配对区间。仅S1为主要比较，其余场景为探索性逐项区间。负值有利于所选策略。')
    h('7 历史预测与下游决策结果')
    h('7.1 预测误差的跨货币差异',2)
    rows=[]
    for currency,c in hist['currencies'].items():
        for model,m in c['forecast_evaluation']['test']['metrics'].items(): rows.append([currency,model,f'{m["qlike"]:.6f}',f'{m["mse"]:.4e}'])
    table(['货币每欧元','模型','测试QLIKE','测试MSE'],rows,[1.1,1.7,1.8,2.05])
    fig('05_historical_volatility.png','图5 测试期因果波动率预测。曲线由当期及过去信息产生，系数未用测试期重新拟合；年化252步是统一约定。来源 historical_forecasts.csv。')
    fig('06_forecast_comparison.png','图6 相对Rolling63的测试QLIKE差。负值表示较低代理损失；点估计排序不是验证选择的替代，也不代表下游经济收益排序。')
    p('USD的Ridge相对Rolling63的QLIKE差约为−0.004521，块长21的95%区间为[−0.027194, 0.013618]，包含零。JPY对应差约为−0.070042，区间[−0.129258, −0.022762]。块长5、21、63下USD均包含零，JPY均低于零；但仍需保留非平稳、代理噪声、多重比较以及忽略模型选择不确定性的限制。')
    fig('08_block_sensitivity.png','图7 Ridge相对Rolling63的QLIKE差对移动块长度的敏感性。三种块长全部展示，每种2000次重采样。不能从这些逐项区间推出全局覆盖。')
    h('7.2 更好的预测是否带来更好的对冲',2)
    rows=[]
    for currency in ['USD','JPY']:
        for strategy in ['rolling63_daily','ridge_daily','selected_daily','selected_every5','selected_threshold']:
            m=metric(currency,strategy);rows.append([currency,strategy.replace('selected_','所选预测 ').replace('_daily',' 每日'),f'{m["rmse"]:.6f}',f'{m["mean_cost"]:.6f}',f'{m["es95"]:.6f}'])
    table(['货币','策略','净RMSE','平均费用','ES95'],rows,[.55,2.1,1.33,1.33,1.34])
    p('JPY的Ridge每日对冲净RMSE为0.419217，Rolling63为0.424266，但两者平方净损失差在3段块长下的95%区间为[−0.017831, 0.007734]。因此，方差代理预测上的优势尚不能转述为明确的下游对冲优势。这种不同目标间的证据落差，是研究发现本身。')
    p('USD的所选预测每日策略平均费用0.114798，阈值策略降至0.096046，但净RMSE由0.317515升至0.337122。JPY费用由0.126016降至0.117697，净RMSE由0.425496变为0.424957，点估计接近。两货币阈值与每日策略的平方损失差区间均包含零，不能据点估计宣布普遍风险改善。')
    fig('07_historical_hedging.png','图8 历史测试的费用与净复制风险。5 bps、每货币60段、起始价格归一为100。竖线0.12为验证预算；JPY每日策略测试平均费用超过此线仍如实展示。')
    p('测试每货币只有60段，最差5%尾部约由3段决定，因此ES估计敏感，不能用六位小数制造可靠性印象。CSV保留完整精度便于复算，表内小数只服务于核对。两条货币路径共享欧元及宏观环境，研究不进行合并样本的显著性宣称。')
    h('8 软件结构与可验证实现')
    table(['层','实现内容','验证依据'],[
        ['数据','官方快照 来源许可 日期质量 SHA256','原始与处理文件hash 缺失重复检查'],
        ['模型','BS Greeks IV CRR MC 路径过程 Ridge EWMA','平价 边界 差分 收敛 未来扰动测试'],
        ['决策','共同premium 固定与阈值规则 验证预算','现金流手算 费用终值恒等式 时间切分'],
        ['证据','JSON CSV SQLite 来源外键与只读查询','输出一致性 SQL聚合与CSV核验'],
        ['交互','定价 敏感度 模拟 历史与能力证据','真实API 请求界限 错误反馈 浏览器流程'],
        ['复现','冻结协议 固定种子 版本与报告构建','测试结果 产物审计 独立复现入口']
    ],[.6,3.15,2.9])
    p('交互工作台在本地运行。参数变更会调用真实计算核心；保存的正式实验与临时交互实验分开标示，避免用户在界面试参数后误认为论文已更新。API限制路径数、步数与总体数组规模，对非有限值、异常JSON及越界输入给出明确错误。服务默认只监听127.0.0.1，无账户、真实交易或资金操作。')
    p('SQLite证据库包含来源表、模拟条件表、预测表与episode表；主键阻止重复观察，外键把每行指向来源CSV和哈希。test_risk视图按货币、策略与成本聚合测试损失；查询脚本使用参数绑定与只读连接。此模块展示数据组织与可审计查询，未把单机研究库描述成大规模生产数据库。')
    h('8.1 正确性测试与产物审计',2)
    count=validation.get('tests_passed')
    p((f'最终记录中共有{count}项自动化测试通过。' if count else '自动化验证的最终数量和运行记录保存在results/validation.json。')+'测试重点是会影响研究结论的性质：看涨看跌平价、Greek差分、无套利界、零波动率与到期处理、反变量独立单位、独立账本手算、费用与财富恒等式、时间边界、未来数据扰动不改变过去预测、验证选择不受测试期修改影响，以及服务输入验证。')
    p('另一个产物审计器核验协议与实现哈希、完整场景覆盖、仅用验证指标所作选择、历史共同premium、账本终值以及SQLite与CSV逐列一致。测试通过只支持已验证范围，不证明不存在所有软件缺陷，更不证明经济假设有效。CI配置随包提供；本地执行和远程CI执行在证据中明确区分。')
    h('8.2 并行计算的实际负结果',2)
    rows=[];medians={}
    for workers in [1,2,4]: medians[workers]=statistics.median(x['elapsed_seconds'] for x in parallel['timings'] if x['workers']==workers)
    for workers in [1,2,4]: rows.append([workers,f'{medians[workers]:.5f}',f'{medians[1]/medians[workers]:.4f}'])
    table(['工作进程数','三次运行中位秒数','相对串行速度'],rows,[1.5,2.6,2.55])
    p(f'基准将200万终端路径分为8个固定SeedSequence子流，共100万个独立反变量配对单位。串行、2进程和4进程使用同一任务量及相同随机块，统计值完全一致，价格{parallel["estimator"]["price"]:.8f}，标准误{parallel["estimator"]["standard_error"]:.8f}。计时纳入启动、计算、回收与合并，重复三次并轮换模式顺序。')
    p('本机上多进程显著慢于向量化串行，说明任务开销超过当前规模的并行收益。项目证明实现了可核验的分块并行与性能测量，不宣称实现集群、GPU计算或取得加速。未来扩大任务、复用进程池等变化应重新按等量工作比较，而不能删掉这一负结果。')
    h('9 十个目标项目的能力并集')
    p('映射依据为2026年10月5日核验的十个项目官方招生或培养页面[P1至P10]。这里取相关能力的并集：保留不同项目独有的关注点，再区分本项目直接形成的证据、部分可迁移的证据和必须由其他经历或材料证明的能力。培养课程、推荐准备和正式入学门槛不是同一类要求；本项目没有声称穷尽所有招生细则，也不声称用一个项目替代整个硕士课程。')
    table(['项目','官方重点的相关部分','本研究连接与边界'],[
        ['CUHK DSBS','统计推断 金融时间序列 Monte Carlo 研究','预测与推断 数值实验 报告 对金融背景直接相关'],
        ['CMU MISM BIDA','分析 决策 软件数据 沟通管理','预算决策 SQL与交互证据 管理协作需另证'],
        ['HKU CS Financial Computing','金融数学 衍生品计算 定量软件','定价库与风险研究 课程准备仍需独立核对'],
        ['Berkeley Analytics','统计 仿真 优化 业务决策','选择流程与成本风险 不宣称全球最优'],
        ['Penn Systems Engineering','系统建模 数据 优化与金融应用','完整数据至决策链 有限候选约束选择'],
        ['Penn DSAI','统计 学习 算法 数据系统 责任','实际Ridge 时间验证 数据审计'],
        ['NTU Data Science','数学 数据准备 建模 解释与应用','完整研究管线 非只做前端展示'],
        ['Imperial ACSE','数值计算 科学编程 并行 研究','树与MC 并行负结果 无集群实证'],
        ['Harvard Health Data Science','数学 编程 统计 健康兴趣','定量方法可迁移 健康领域兴趣另证'],
        ['Yale MPH Biostatistics','数学 统计 公共卫生与协作','推断与不确定性可迁移 非临床或公共卫生研究']
    ],[1.5,2.15,3.0])
    p('因此并集不意味着每个能力都被这个金融项目完全实现。健康领域问题定义、公共卫生动机、真实团队合作、正式学分成绩和个人掌握程度必须保留为外部证据。与其他已完成项目连接时，可以说明方法如何迁移，但不能把金融汇率分析改名为健康数据分析。Yale条目明确为MPH Biostatistics，Harvard条目为Health Data Science SM；不混用其他学位的要求。')
    rows=[]
    for c in union['competencies']:
        cid=c['id'];status='直接证据' if cid not in ['C13','C14','C15','C16','C17','C18'] else ('部分证据' if cid in ['C13','C14'] else '需要外部证据')
        evidence=c.get('evidence',[]);evidence='；'.join(evidence) if isinstance(evidence,list) else str(evidence)
        rows.append([cid,c['name'],status,evidence])
    table(['编号','能力维度','覆盖性质','定位证据'],rows[:12],[.45,1.65,.85,3.7])
    p('保留部分覆盖与外部证据，避免将能力并集缩减为本项目恰好能做的内容。')
    table(['编号','能力维度','覆盖性质','定位证据'],rows[12:],[.45,1.65,.85,3.7])
    p('上述定位用于向读者解释研究内容，不是招生打分表。课程证书、报告页数和代码行数都不能自动转化为录取概率。个人使用本成果时，应能解释数据时序、关键公式、失败结果及自己真实完成的修改；无法解释的部分应作为学习计划，而非已有能力陈述。')
    h('10 局限与可检验的后续方向')
    p('第一，历史价格是信息参考价。没有实际可成交买卖价、滑点、价差、发布延迟建模、保证金或冲击成本；r=q=0也没有体现真实外汇双利率。结果应称为历史路径上的假设复制实验，而不是实盘收益回测。即使模型P&L为正，也不足以证明可执行套利。')
    p('第二，平方收益代理噪声大，条件均值并非严格为零；单步方差外推到剩余期权期限忽略了期限结构。Ridge校正、固定特征和长期冻结系数都可能在结构变化下失效。更复杂模型需要新的验证设计与未触碰测试数据，不能仅因当前测试结果不好而无记录地迭代到满意为止。')
    p('第三，区间存在假设。模拟区间条件于生成机制；移动块区间依赖弱依赖和近似平稳假设，未包含模型选择不确定性。仅有60段的尾部风险尤其脆弱。主要比较以外的探索性结果需要重复证据，不应以多次尝试中最显著的一项作为核心成功故事。')
    p('第四，有限候选策略的验证选择不等于全局优化。平均费用预算不控制最差路径，融资和偿付假设也较宽松。后续可研究ES或机会约束，但应先定义可行性、数据需求与求解误差，再决定是否值得增加算法复杂度。')
    p('第五，软件为可复现研究工具，未经过生产安全审计、长期负载验证或真实用户研究。本地并行性能结果也不能外推至集群。真实团队协作、业务采纳和个人独立掌握不由自动测试证书替代。')
    table(['下一项研究','所需新增证据','预先明确的检验'],[
        ['期限匹配的方差预测','新的未观察时间窗或独立市场','与单步外推在相同episode上比较'],
        ['真实期权定价与对冲','带时间戳报价 双利率 执行成本','分离价格误差 复制误差与执行误差'],
        ['尾部预算约束','足够尾部样本与明确风险容忍度','样本外可行率与尾部损失而非只看均值'],
        ['长期部署与协作','真实使用记录 代码审查 权限设计','用户错误恢复 维护成本与复现一致性']
    ],[1.5,2.4,2.75])
    h('11 结论')
    p('本研究的核心结果是区分并连接三个层次：公式和程序能否正确计算，模型能否在未参与选择的数据上准确估计，以及该估计能否改善有成本约束的最终决策。模拟支持模型匹配、低费用下频繁调仓的条件性优势；历史实验同时展示预测改进未必传到对冲、费用下降可能伴随风险上升。保留这些负结果，使成果能够支撑对方法、假设和证据的讨论。')
    p('配套项目将这些问题落实为可运行代码、冻结数据、可复算账本、统计区间、可查询证据和交互界面。它适合证明量化研究与数据软件实践的连接，也保留了健康领域经验、正式先修与个人贡献等必须独立证明的边界。其价值在于每一项结论都有可追踪的计算依据，而不是功能或篇幅的堆积。')
    h('参考资料')
    p('文献支持方法出处；院校页面支持能力映射。所有网络资料核验日期为2026年10月5日，课程或招生政策可能在后续申请周期更新。')
    for ref in refs:
        par=p(f'[{ref["id"]}] {ref["citation"]} ');link(par,'来源页面',ref['url'])
        par.paragraph_format.keep_with_next=False
    h('附录A 文件与复现步骤')
    p('在解压后的项目根目录创建Python 3.11或更高版本虚拟环境，安装pyproject.toml声明的研究依赖及dev、report可选依赖。requirements-reproduced.txt记录本次运行使用的直接依赖版本；平台计时差异和合理浮点误差不视为结果造假。')
    table(['顺序','命令','输出或作用'],[
        ['1','python -m pip install -e ".[dev,report]"','安装本地项目及测试和报告依赖'],
        ['2','python -m pytest','运行核心 历史 产物与接口测试'],
        ['3','python scripts/run_simulations.py','重现正式模拟及数值基准'],
        ['4','python scripts/run_historical.py','在冻结快照上重现历史结果'],
        ['5','python scripts/benchmark_parallel.py','同量分块的串行和进程对照'],
        ['6','python scripts/build_research_store.py','生成SQLite证据快照'],
        ['7','python scripts/verify_artifacts.py','交叉核验协议 CSV JSON与数据库'],
        ['8','python scripts/build_figures.py','从正式结果生成8幅图'],
        ['9','python scripts/build_report.py','从正式结果生成本报告'],
        ['10','python -m risklab.server --port 8870','打开本地研究工作台']
    ],[.4,3.55,2.7])
    p('浏览器地址为http://127.0.0.1:8870。快速模拟使用--quick，输出到results/smoke，不覆盖正式论文结果。fetch_ecb.py默认复用归档快照；只有显式--refresh才下载新快照，更新会改变来源证据并要求重新生成下游成果。后续算法改变须记录版本与已查看测试集的事实。')
    p('正式模拟协议SHA256 '+sim['protocol_sha256'])
    p('处理后ECB数据SHA256 '+hist['source']['files']['ecb_usd_jpy.csv']['sha256'])
    h('附录B 贡献与结果追溯')
    p('results目录提供252条模拟条件汇总、21项配对比较、10752行历史预测与4536个历史策略费用记录。后两者不能误作独立样本数量：预测包含训练期拟合值，对冲测试每货币仅60段。docs/evidence_index.md逐项列出报告图表与原始结果对应关系；数据与验证哈希用于复核来源和实现。')
    p('贡献边界与个人学习记录说明见docs/contribution_record.md；研究理解检查见docs/research_defense.md。开发和报告使用实质性AI协助，没有据此推定独立研究、导师指导、发表、实习时长或团队管理经历。')
    output=OUT/'金融数学风险研究_论文式报告.docx';doc.save(output)
    print(json.dumps({'report':str(output),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes)},ensure_ascii=False))


if __name__=='__main__':main()
