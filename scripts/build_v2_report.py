"""Generate the V2 research manuscript and publication assessment from actual results."""
from __future__ import annotations
import json,re,hashlib
from pathlib import Path
from statistics import mean
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
TITLE='期限匹配与有限样本选择下的期权对冲研究'
def load(name):return json.loads((ROOT/name).read_text(encoding='utf-8'))
def fmt(x):return f'{x:.6f}'
def interval(r):return f'[{r["ci_low"]:.6f}, {r["ci_high"]:.6f}]'

def main():
    hist=load('results/v2_historical_summary.json');sim=load('results/v2_selection_summary.json');lp=load('results/v2_convex_summary.json')
    validation=load('results/validation.json');routes=load('docs/publication_routes.json');literature=load('docs/v2_literature.json')
    primary=sim['primary_comparison'];retro=[r for r in hist['primary_comparisons'] if r['primary_block'] and r['period']=='retrospective']
    additional=[r for r in hist['primary_comparisons'] if r['primary_block'] and r['period']=='additional_2026']
    matched=next(r for r in lp['contrasts'] if r['scenario']=='matched_heston' and r['metric']=='es90')
    shifted=next(r for r in lp['contrasts'] if r['scenario']=='shifted_heston' and r['metric']=='es90')
    maxvio=max(r['position_violation_fraction'] for r in lp['test_rows'] if r['method']=='convex_cvar')
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.73);sec.bottom_margin=Inches(.70);sec.left_margin=sec.right_margin=Inches(.80)
    for name in ['Normal','Body Text','Caption','Title','Subtitle','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ['Normal','Body Text','Caption'] else '微软雅黑')
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.2;doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    for name,size in [('Title',23),('Subtitle',12),('Heading 1',15),('Heading 2',12),('Caption',9)]:doc.styles[name].font.size=Pt(size)
    doc.styles['Subtitle'].font.italic=False
    for name in ['Heading 1','Heading 2']:
        st=doc.styles[name];st.font.bold=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
    for st in doc.styles:
        for b in list(st._element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title=TITLE;doc.core_properties.subject='FinMath Risk Lab V2 研究报告与发表可行性评估'
    doc.core_properties.author='Project initiated by Youming Yang; material AI assistance disclosed'
    def p(text,style=None):return doc.add_paragraph(text,style)
    def h(text,level=1):return doc.add_heading(text,level)
    def page(title):
        par=h(title)
        if title.startswith(('1 ','16 ','参考文献')):par.paragraph_format.page_break_before=True
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Inches(w)
        for c,txt in zip(t.rows[0].cells,headers):c.text=str(txt)
        for row in rows:
            for c,value in zip(t.add_row().cells,row):c.text=str(value)
        for i,row in enumerate(t.rows):
            tr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');tr.append(no)
            if i==0:tr.append(OxmlElement('w:tblHeader'))
            for j,(c,w) in enumerate(zip(row.cells,widths)):
                c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders);pad=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');pad.append(e)
                pr.append(pad)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EEF1');pr.append(shade)
                for par in c.paragraphs:
                    par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2);par.paragraph_format.line_spacing=1.08
                    if i==0:par.paragraph_format.keep_with_next=True
                    for r in par.runs:r.font.size=Pt(9);r.bold=(i==0)
        p('').paragraph_format.space_after=Pt(1)
    def fig(name,caption):
        par=p('');par.paragraph_format.keep_with_next=True
        im=par.add_run().add_picture(str(ROOT/'figures'/f'{name}.png'),width=Inches(6.55));im._inline.docPr.set('descr',caption)
        p(caption,'Caption')
    def eq(text):
        par=p('');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        mathp=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);m.append(r);mathp.append(m);par._p.append(mathp)
        par.paragraph_format.space_after=Pt(7)
    def link(par,label,url):
        node=OxmlElement('w:hyperlink');node.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);node.append(r);par._p.append(node)
    p(TITLE,'Title');p('金融数学计算研究与论文发展评估','Subtitle')
    p('FinMath Risk Lab V2   2026年10月6日\n项目发起人  杨又铭   Python研究库与可复现实验配套报告')
    h('摘要')
    p('本研究考察期限匹配、有限验证样本和风险目标如何影响含比例交易成本的假设期权对冲。研究将真实历史参考汇率的年度滚动评估、完整选择流程的重复模拟，以及可解释CVaR线性规划学习器放入可审计的自融资账本。重点是解释方法的适用条件和失效机制，而不是以更复杂模型替代下游风险检验。')
    p(f'六种货币共{hist["counts"]["currency_year_folds"]}个年度折中，2015至2025年期限Ridge相对单步外推的六个主要区间均包含零。模拟的预先指定主比较为匹配Heston、64条验证路径、5 bps费用；联合MSE选择相对预测优先选择的测试MSE差为{fmt(primary["difference"])}，95%区间{interval(primary)}，独立完整流程重复24次，没有形成改善证据。')
    p(f'有限特征CVaR学习器通过24次线性规划和独立账本核验。在八次完整重复中，其相对验证选择带状策略的ES90差在匹配场景为{fmt(matched["difference"])}，在分布变化场景为{fmt(shifted["difference"])}。局部训练目标改善不能保证分布外风险改善，样本内仓位约束也不保证未来状态约束。所有结果保留原始方向，不以测试集重新调参。')
    h('Abstract')
    p('We study maturity alignment and finite-sample procedure selection for hypothetical option hedging with proportional costs. Three linked experiments combine purged yearly evaluation on six ECB reference exchange-rate series, repeated calibration-validation-test simulations, and a finite-feature CVaR hedger solved by linear programming. Independent cash accounting, solver certificates and explicit information boundaries make the results auditable.')
    p(f'The primary simulated contrast, joint-MSE selection minus forecast-first selection under matched Heston dynamics, yields a test-MSE difference of {primary["difference"]:.6f} with a 95% interval {interval(primary)} across 24 complete repetitions. Historical maturity alignment does not establish a consistent advantage. An interpretable CVaR learner improves tail loss under matched dynamics but deteriorates after a distribution change. These conditional findings motivate a reproducible methods-evaluation manuscript; they do not establish a novel hedging algorithm, executable market performance, or guaranteed publication.')
    p('关键词  金融数学  期限匹配  选择不确定性  条件风险价值  模型风险  可复现研究')

    page('1 研究问题与已有工作的关系')
    p('课程中的无套利和复制思想把价格、概率与交易策略连在一起；数据驱动实现又引入估计误差和模型选择。一个方差预测器在统计损失上更好，不意味着它输入delta策略后一定降低净终端损失。本研究进一步追问：剩余期限是否匹配，验证样本能否稳定选出策略，以及直接学习风险目标后能否承受分布变化。')
    p('金融决策目标与预测目标的区别至少可追溯到Bengio的工作[L01]。无交易区间及比例成本下交易到边界也已有Whalley与Wilmott的研究[L02–L03]。Deep Hedging和No-Transaction Band Network已系统研究摩擦市场中的风险学习[L05–L06]。因此，本报告将这些方法视为基线和设计依据。')
    table(['问题','本次可核验证据','与已有工作的边界'],[
        ['期限错配','1至21步直接目标与单步外推的年度比较','不提出新方差模型'],
        ['选择误差','24次独立校准与选择，三个验证规模','不同于只给固定策略大量测试路径'],
        ['风险目标','MSE与ES90共同预算比较','直接优化金融目标并非首创'],
        ['分布变化','训练验证不变，仅测试与参考机制变化','不把匹配内改善外推到全部市场'],
        ['可解释优化','仿射仓位与CVaR线性规划及证书','已知凸构造，不主张理论创新']],[1.0,2.65,2.90])
    p('2026年的近邻文献进一步缩小了可声称的空白：Rethinking Synthetic Scenario Realism讨论生成机制与对冲器的兼容性[L10]；结构先验带状网络研究将随机控制结构融入学习策略[L09]。单独增加Heston、CVaR、分布漂移或可解释特征，都不足以构成原创性。')
    p('本研究目前较清楚的定位，是把期限、验证规模和目标选择放在同一信息约束与现金流终点下，记录可重复的条件性和负结果。是否形成发表贡献，仍需证明这组比较回答了未被近邻文献充分回答的问题。定向文献检索不是穷尽性系统综述，也不是原创性认证。')

    page('2 数学对象与自融资评价')
    p('统一研究卖出一份欧式看涨期权后的复制误差。记离散价格为S₀至Sₙ，执行价K，期权收益H=max(Sₙ−K,0)，共同初始权利金P。持仓δₜ在观察时点t的信息后确定，并用于随后一段价格变化。V2实验使用r=q=0、21个观察间隔，每个episode重新归一化S₀=K=100。')
    eq('δ₋₁ = 0    δₙ = 0    C = c Σₜ₌₀ⁿ Sₜ |δₜ − δₜ₋₁|')
    eq('Π = P + Σₜ₌₀ⁿ⁻¹ δₜ (Sₜ₊₁ − Sₜ) − C − H    L = −Π')
    p('成本同时计入建仓、中间调仓和到期清仓。一般账本函数还支持现金利息与成本终值化；V2线性规划与三项研究明确使用零利率，因此上述简式与逐笔账本一致。gross P&L是同一路径与持仓去除成本后的损益，不是另一条免费交易策略。')
    eq('MSE(Π) = mean(Π²) = Varₙ(Π) + mean(Π)²')
    p('分解中的方差使用ddof=0；样本标准差的无偏修正不能直接放进这一恒等式。均方损失同时惩罚波动和均值偏移。共同保费在不同生成机制下可能失配，因此均值损失中包含价格失配，不能全部解释为对冲技能。')
    eq('ESα(L) = minη { η + mean[max(L − η, 0)] / (1 − α) }')
    p('本版统一报告ES90，即卖方损失最坏10%经验质量的均值。排序后在尾部边界用部分权重，避免样本数不能整除时偷偷增加尾部概率。ES可能为负，这表示选定尾部仍获利；代码不把负值截为零。2026每币仅9个episode，尾部质量0.9，对ES解释尤其脆弱。')
    h('物理概率与定价概率',2)
    p('历史模型估计实际参考收益在物理概率下的条件二阶矩，随后把其平方根输入Black–Scholes delta[R9]作为控制启发式。该输入不是市场隐含波动率，也没有证明完成了P到Q的风险溢价转换。ECB参考汇率[D1]不是成交价，本研究未取得真实外汇期权报价或双边利率，因此不把实验称作可执行外汇期权定价或盈利回测。')

    page('3 可解释尾部风险学习的凸构造')
    p('为超越有限网格挑选，本版加入直接学习仓位的优化器。基准delta来自常数年化波动率0.20的Black–Scholes公式，四个固定因果特征依次为常数、截断对数价内程度、剩余期限比例及4Δ(1−Δ)。这些特征只依赖当前价格与时间。')
    eq('δᵢₜ(θ) = Δᵢₜ + φᵢₜᵀ θ    φᵢₜ ∈ R⁴    |θⱼ| ≤ 0.25')
    p('固定特征中的截断保留了仓位对参数θ的仿射性。优化后不对动作另行clip；否则得到另一策略，线性规划证书与实际账本便可能不再对应。比例费用是仿射交易量绝对值之和，因而单路径损失关于θ为凸的分段线性函数。')
    eq('Lᵢ(θ) = Hᵢ − P − Σₜ δᵢₜ(θ) ΔSᵢₜ + c Σₜ Sᵢₜ |Δδᵢₜ(θ)|')
    p('采用Rockafellar与Uryasev的CVaR表示[R8]，引入每笔绝对交易上界aᵢₜ、尾部超额uᵢ、阈值η与系数绝对值上界zⱼ。求解目标和约束如下。')
    eq('min  η + Σᵢ uᵢ / [n(1−α)] + λ Σⱼ zⱼ')
    eq('aᵢₜ ≥ Δδᵢₜ    aᵢₜ ≥ −Δδᵢₜ    zⱼ ≥ θⱼ    zⱼ ≥ −θⱼ')
    eq('uᵢ ≥ Hᵢ − P − Σₜ δᵢₜ ΔSᵢₜ + c Σₜ Sᵢₜ aᵢₜ − η    uᵢ ≥ 0')
    eq('meanᵢ[c Σₜ Sᵢₜ aᵢₜ] ≤ B    0 ≤ δᵢₜ ≤ 1')
    p('所有表达式都对优化变量线性。绝对交易上界扩大只会增加或维持尾部风险并消耗预算，所以总能选取等于真实绝对交易量的最优表示。部分非尾部路径的松弛变量可能不紧，这不影响存在紧表示；执行后还必须用真实交易量独立重算目标。这里得到的是已知有限特征和训练样本下的凸优化问题，不是所有可适应策略空间的全局最优对冲定理。')

    page('4 优化证书与实验信息边界')
    p('每个完整重复使用192条匹配Heston训练路径、128条独立验证路径和每场景4096条测试路径。对λ=0、0.01、0.1分别训练；验证集在平均成本预算0.18下选择ES90最小者。同一验证集也在带宽0.05与0.10之间选择带状基线。八组完整重复分别使用独立随机种子。')
    table(['核验项','本次执行值','含义'],[
        ['成功拟合',f'{lp["audit"]["fits"]} / {lp["audit"]["fits"]}','所有候选均正常求解'],
        ['每次变量数',lp['fits'][0]['variables'],'4个参数及绝对值与尾部辅助变量'],
        ['每次不等式数',lp['fits'][0]['constraints'],'仓位 预算 绝对交易与尾部约束'],
        ['最大绝对对偶间隙',f'{lp["audit"]["max_abs_duality_gap"]:.3e}','求解器原始与对偶目标比较'],
        ['最大独立目标误差',f'{lp["audit"]["max_objective_reconciliation_gap"]:.3e}','真实账本ES加L1项对照LP目标'],
        ['最大不等式违反',f'{lp["audit"]["max_constraint_violation"]:.3e}','数值容差内的训练可行性']],[1.65,1.25,3.65])
    p('审计还保存一阶驻点残差、实际训练成本、训练仓位极值及每次系数。独立测试验证了损失凸性、不可能预算下返回不可行状态、未来价格扰动不改变过去特征，以及优化目标与现金流的一致性。求解成功只说明有限训练问题被数值求解，不能替代样本外验证。')
    p('仓位0至1与平均成本预算都是训练样本约束。由于模型对状态的响应包含未经全域限制的基准delta加仿射修正，在新状态上仍可能越界。报告保留真实越界比例和仓位极值，不用自动修正掩盖这种泛化缺口。若未来用于实际交易，还需要独立定义可执行约束与重新评价策略；本项目目前是研究原型。')
    p('实现审计曾发现每日基线的仓位极值字段使用了理论上下界，而非样本观测值。已改为重新计算同一因果delta序列，并完整重跑凸优化实验。修复未改变损益、风险或比较结论；最终结果哈希对应修复后的代码。')

    page('5 历史期限预测与年度滚动协议')
    p('历史研究使用ECB归档参考汇率中的USD、JPY、GBP、CHF、CAD、AUD六列，单位均为一欧元对应多少该货币。训练起点为2005年，2015至2025年逐年作为外层评估，2026年1月1日至9月30日为追加时期。每个外层年度仅使用之前两年作内层验证，模型和策略每年重新选择。')
    eq('yₜ⁽ʰ⁾ = (1/h) Σⱼ₌₁ʰ rₜ₊ⱼ²    h ∈ {1,…,21}')
    p('期限Ridge分别估计1至21个观察间隔的平均未来平方收益。每个决策时点按期权剩余步数调用相应模型；flat模型只估计下一步并外推。七项因果特征沿用基线；标准化、对数目标反变换的均值比例修正均只由训练样本拟合。')
    table(['模型','输入与估计','剩余期限处理'],[
        ['rolling63','63个过去平方收益均值','对全部剩余期限保持同一估计'],
        ['ewma97','过去平方收益的指数加权','同上'],
        ['ridge_flat','训练历史上的单步log目标Ridge','把h=1预测当作全期限二阶矩'],
        ['ridge_term','分别拟合h=1至21的Ridge','按当前剩余h读取直接预测']],[1.2,2.95,2.40])
    p('清除跨界标签按目标结束日期执行，而不是只检查特征日期。初始训练标签必须在内层验证开始前结束；内层21步评分窗口必须完全落在验证期。α从0.1、1、10、100中按21步QLIKE选择，之后使用外层年度前已完成的全部标签重新拟合。QLIKE与不完美波动代理下的预测评价有已有研究[R10]；引用该工作不意味着本实验自动满足其全部理论条件。测试期不会参与参数、scaler或修正因子的估计。')
    p('各模型共用episode起点rolling63对应的保费，策略包括每日、每五步、边界带宽0.05和0.10。在5 bps下，forecast-first先选QLIKE模型再选预算内MSE策略；joint-MSE与joint-ES90分别直接在模型和策略组合中选取。预算为0.12。选择结果再在0、5、10 bps上评价，不按外层成绩挑选费用。')
    p('若预算内候选不存在，程序明确记录最低成本回退。完整结果包含每次决策真正使用的预测，以及内层选择、最终重拟合的日期边界。未来数据扰动测试用于核验过去决策不被改写。')

    page('6 回顾期没有形成普遍的期限改进证据')
    p('主要历史比较固定为ridge_term减ridge_flat、带宽0.05、5 bps的episode平方净损失差。每种货币有132个不重叠episode；这些时间段仍可能相关，因此以连续episode的移动块bootstrap评价，不把六种共同受市场冲击的货币当作六组独立重复。')
    fig('v2_history_retrospective','图1  回顾期不同块长的95%逐项区间。各小图使用自己的横轴尺度，负值有利于期限模型；CHF的尺度明显更大。')
    table(['货币','平均差','块长3区间','Holm p'],[[r['currency'],fmt(r['difference']),interval(r),f'{r["holm_six_currency_p"]:.3f}'] for r in retro],[.65,1.20,3.80,.90])
    p('所有主要区间包含零，六币Holm调整后的p值均为1，不能据此声称期限建模普遍降低对冲风险。置信区间并非同时覆盖区间；Holm只作用于块长3、同一时期的六项检验。CHF的大尺度同时提醒读者，极端历史episode会显著影响均方损失与相关重采样。')
    p('回顾期是探索性证据。V1已经查看过USD和JPY在2021至2025年的结果，即使V2增加了年度重拟合，也不能把这些历史年份重新称为未看过的确认集。每个fold的信息顺序正确，与研究设计完全未受既往结果影响，是两个不同问题。')

    page('7 追加时期与统计结论的脆弱性')
    p('2026追加时期来自此前归档但未在V1评分的数据，在V2协议冻结后首次运行。它提供额外检查，却不是外部预注册试验：该时期短，研究者已掌握早期结果，且仍只有同一类参考价格。每币仅9个完整episode，不足以稳定估计尾部和时间依赖。')
    fig('v2_history_additional_2026','图2  2026年1至9月追加时期。每币n=9；块长1、3、6的比较是敏感性检查，不是三个独立实验。')
    aud=next(r for r in additional if r['currency']=='AUD')
    p(f'AUD的term减flat平均差为{fmt(aud["difference"])}，块长3的95%区间为{interval(aud)}，六币Holm p为{aud["holm_six_currency_p"]:.4f}，方向是不利于期限模型。该信号对块长和极少episode的依赖较强，因此保留为脆弱的探索性负证据，不作稳健显著性结论。')
    p('实现采用4000次重叠非循环移动块抽样，并在中心化零假设下计算近似双侧p值，加入有限抽样的加一修正。相关序列块bootstrap的经典方法来源见[R11]。百分位区间与零假设p值来自不同近似，不保证在有限样本中完全对偶。该方法还依赖近似平稳与弱依赖；突发制度变化会削弱这些近似。')
    p('历史重采样以已得到的年度流程损失为对象，并没有在每个bootstrap内部重新训练所有模型和选择全部策略。它不能完全消除训练不确定性、研究者选择不确定性或长期非平稳性。为此，另设可控模拟中的完整流程重复，直接观察有限校准与验证对最终表现的影响。')

    page('8 完整选择流程的重复模拟')
    p('每次重复先独立生成60个校准收益，估计年化波动率；候选预测为该估计的0.75、1.0和1.25倍。每个预测与四种交易规则组合，形成12个共同候选。三个选择器使用完全相同的候选、路径、保费和预算，避免把信息优势误称为方法优势。')
    table(['因素','冻结设置'],[
        ['价格与期限','S₀=K=100，21步，T=21/252，r=0，μ=0'],
        ['场景','匹配GBM、匹配Heston、Heston变化、跳跃变化'],
        ['验证规模','16、64、256；同一重复内使用同一256条池的嵌套前缀'],
        ['独立重复','每场景24次完整校准与验证流程'],
        ['测试与参考','每重复2048条测试，另有16384条独立参考路径'],
        ['费用与预算','5 bps预算0.18；20 bps预算0.72'],
        ['保费','统一按σ=0.20计算，跨候选和场景保持一致'],
        ['评分','净MSE、方差、均值平方、ES90、成本和预算越界']],[1.3,5.25])
    p('GBM场景的校准与验证使用GBM。其他三个场景均在匹配Heston上校准和验证，只有测试及独立参考路径分别改变为匹配Heston、高波动Heston或补偿跳跃机制。因此场景变化反映冻结选择程序面对未知机制的表现；选择器不能提前看到变化后的样本。')
    p('预测优先按平均未来平方收益的QLIKE选波动率，再在该预测下按预算内MSE选交易规则；联合MSE和联合ES90在全部组合中选择。无可行候选时保留明确的最低成本回退。n=16时经验ES90只对应1.6个尾部观测质量，风险目标本身也会有较强估计噪声。')
    p(f'正式执行产生{sim["counts"]["candidate_rows"]:,}条候选记录与{sim["counts"]["selected_rows"]:,}条选择记录。区间的独立单位是24次完整重复，不能把同一个训练选择程序下的2048条测试路径当作2048次独立建模实验。不同验证规模、费用和选择器共享随机结构，比较均保留配对关系。')

    page('9 主要模拟结果与验证规模')
    p(f'唯一预先指定主比较为匹配Heston、64条验证路径和5 bps费用。联合MSE选择相对预测优先的测试MSE差为{fmt(primary["difference"])}，95%配对t区间{interval(primary)}，n={primary["n"]}。点估计略差且区间跨零，不能把联合优化描述为已证实更优。')
    fig('v2_selection_size','图3  验证规模与最终净MSE。各面板纵轴不同；线段只连接三个已执行规模，不代表连续响应函数。图中其余比较均属探索。')
    mainrows=[r for r in sim['aggregates'] if r['scenario']=='heston' and r['validation_n']==64 and r['fee_bps']==5]
    table(['选择器','平均MSE','平均ES90','平均成本'],[[r['selector'],f'{r["mean_test_mse"]:.5f}',f'{r["mean_test_es90"]:.5f}',f'{r["mean_test_cost"]:.5f}'] for r in mainrows],[2.05,1.45,1.50,1.55])
    p('更多验证样本没有令每个场景、目标和随机重复都单调改善。选择误差、目标不一致、候选类不足与分布变化可以共同主导最终结果。图中不按哪个格子最显著来重新定义主要发现；其余逐项区间没有整体多重比较保证。')

    page('10 风险目标之间的取舍与选择稳定性')
    fig('v2_selection_tradeoff','图4  匹配Heston主设定中的MSE分解与风险成本位置。这里展示重复均值；点之间的可见差距不自动构成显著差异。')
    first=next(r for r in mainrows if r['selector']=='forecast_first');tail=next(r for r in mainrows if r['selector']=='joint_es90')
    p(f'在主设定中，ES90选择的平均尾部损失为{tail["mean_test_es90"]:.5f}，预测优先为{first["mean_test_es90"]:.5f}；但前者平均MSE为{tail["mean_test_mse"]:.5f}，高于后者的{first["mean_test_mse"]:.5f}。因此应说明目标取舍，而不是挑选一种有利指标宣布全面改善。')
    p('每个设定同时保存候选选择频率、选择熵、最常选择比例和样本外预算违反比例。选择集中可能表示稳定，也可能表示候选很少或模型偏差长期一致；高选择熵可能来自有限样本噪声，也可能是候选表现接近。这些描述量本身不是通用稳定性定理。')
    p('独立参考集用于构造有限候选类的信息基准：在参考平均成本预算内选择参考MSE最小的候选，并在同一参考路径上评价实际选择与基准的差。参考样本虽更大仍有Monte Carlo误差，其最小值也有选择乐观偏差。该“oracle”只能回顾地读取参考数据，不能作为可部署选择器。')
    p('若实际选择在参考样本上不满足预算，它与预算内参考最优的差可能为负；此时不能称为同一可行集内必然非负的遗憾。文件将其记录为带符号参考差，并同时显示可行性。研究只描述本次有限网格与冻结预算下的行为，不声称获得所有决策方法的性能下界。')
    p('三个验证规模嵌套在同一池中，两档费用也共享路径，这提高了配对比较效率，但不会增加独立重复数。对于尾部目标及制度变化，更可靠的后续确认需要额外冻结实验，而不是把本次所有格子视为独立证据累加。')

    page('11 随机波动率数值方案与步长敏感性')
    eq('dS = μS dt + √v S dWˢ    dv = κ(θ − v)dt + ξ√v dWᵛ')
    eq('corr(dWˢ,dWᵛ) = ρ    v⁺ = max(v,0)')
    p('匹配Heston使用v₀=θ=0.04、κ=2、ξ=0.30、ρ=−0.70；变化场景使用v₀=θ=0.09、κ=1、ξ=0.45、ρ=−0.70。按照full-truncation思路[L11]，原始方差状态可短暂为负，但漂移和扩散使用其正部；价格使用旧时点正部方差进行log-Euler更新。潜在方差只用于生成和诊断，不提供给可部署候选。')
    fig('v2_discretization','图5  每观察间隔1、2、4、8、16个子步的敏感性。每点32768条独立路径；虚线为方差一阶矩理论值，误差条表示Monte Carlo误差。')
    p('正式研究每个观察间隔使用4个子步，交易仍只在21个观察点之间发生。更细子步改善的是生成近似，不是无成本增加交易机会。诊断保留原始方差负更新比例、曾出现负值的路径数量、最小原始方差和Feller条件信息。')
    p('精度敏感性采用不同独立路径集合，并非共享Brownian路径的强误差收敛实验。均值偏差和风险变化同时包含时间离散误差与Monte Carlo误差，不能宣称误差严格单调、阶数已获证明或生成器已等同精确Heston采样。当前结果足以暴露粗网格影响，仍不足以取代更高阶方案的专门数值研究。')

    page('12 直接风险学习的条件性改善')
    fig('v2_convex_contrasts','图6  CVaR学习器减验证选择带状基线。各指标负值分别表示尾部损失、MSE或成本更低；仅八个完整重复，均为探索性逐项95%区间。')
    table(['测试机制','指标','平均差','95%区间'],[[r['scenario'].replace('_heston',''),r['metric'],fmt(r['difference']),interval(r)] for r in lp['contrasts']],[1.05,.85,1.05,3.60])
    p(f'匹配场景ES90差为{fmt(matched["difference"])}，区间{interval(matched)}，支持在这组有限参数和特征下的条件性改善。平均成本增加，因此该结果不是同时更便宜、更低风险。它展示直接学习目标与优化证书可以结合，但没有建立对全部交易成本和市场机制的支配。')
    p(f'变化场景ES90差反向为{fmt(shifted["difference"])}，区间{interval(shifted)}。同一学习器在新的高波动机制中损失更差，说明训练分布拟合与尾部目标之间存在脆弱联系。不能用匹配场景的正结果覆盖这一反例。')

    page('13 分布外约束与模型风险')
    fig('v2_convex_tradeoff','图7  每个点代表一次完整重复。每日delta、验证选择带状策略和CVaR学习器使用相同测试价格与保费；所有点均为实际执行值。')
    p(f'CVaR学习器在最不利一组测试中，有{maxvio*100:.4f}%的决策状态超出训练时规定的0至1仓位区间。这个百分比以路径乘决策时点为分母，不是发生越界的路径百分比。越界数据被保留在结果中；将其自动clip后仍引用原LP证书会混淆两个策略。')
    p('分布变化的影响包括更宽的状态分布、更高的波动和共同保费失配。实验不能识别其中每个因素的独立因果贡献，因此将其称作指定情景下的模型风险，不称作对现实制度变化的因果发现。八次重复也不足以支持非常精细的尾部概率估计。')
    p('当前特征只有四维，成本结构为比例费用，没有固定费用、市场冲击、借券、保证金或多资产交叉影响。LP数值全局最优只适用于这一受限经验问题。加入深度模型、状态全域限制或分布稳健目标可能改变权衡，但必须重新冻结假设与预算，并使用独立评价。')
    p('本版没有训练神经Deep Hedging或No-Transaction Band Network比较器。因而不能声称超过这些方法，也不能将“可解释”直接等同于更高性能。实际科研下一步应由待回答问题决定：若主张学习方法创新，强学习基线不可缺少；若定位可复现方法评价，则应把有限样本失效边界和复现价值讲清楚。')

    page('14 复现工程与验证证据')
    p('配套项目以相同Python研究核心驱动结果和本地工作台。V1保留解析定价、Greeks、隐含波动率、CRR树、Monte Carlo、历史预测、账本及SQLite证据；V2增加期限滚动、随机波动率、重复选择和CVaR优化。新结果文件均带v2前缀，原始V1交付另行保留。')
    table(['证据','实际保存与检查'],[
        ['冻结协议','docs/v2_protocol.json；结果记录协议SHA256'],
        ['来源与时序','ECB快照哈希、六币处理说明、每fold边界与真实决策预测'],
        ['逐项结果','历史123552行、选择1728行、凸优化48个测试汇总'],
        ['优化核验','24个拟合证书、独立账本目标、探索性区间重算'],
        ['软件验证',f'{validation["tests_passed"]}项Python测试通过；具体环境和日志随包'],
        ['交互验证','桌面与移动界面、真实结果、缺失与无效证据状态检查'],
        ['发布完整性','源码统计、文件SHA256清单、ZIP CRC与清单逐项对照']],[1.25,5.30])
    p('在源码目录安装依赖后，运行 scripts/reproduce_v2.py 可依次重现三组V2数值实验和核验；使用 --report 生成图表和文档。默认使用归档数据，不需要重新下载。代码、源数据和固定种子可复核，但依赖版本、数值库和平台会影响极小数值差异与耗时，不承诺逐字节跨平台相同。')
    p('本地服务默认监听127.0.0.1:8872。研究深化页面从结果文件读取实际数字，支持时期与货币选择，同时显示证据缺失和大小限制错误。网站只是查看与解释入口，数值方法可脱离浏览器独立运行；它不连接交易账户。')
    p('Python测试覆盖数学恒等式、时间信息边界、账本成本、LP目标、求解失败、bootstrap与结果哈希。通过软件测试意味着这些明确断言在记录环境中成立，不等于经济假设正确、未来泛化可靠或论文已获同行认可。文档由执行结果生成，且另行渲染核验版式。')

    page('15 研究贡献的合理边界')
    p('本版比V1增加了真正的方法深度：期限目标与标签终点清除、训练和选择的重复不确定性、随机波动率数值诊断，以及带独立证书的直接风险学习。它已超出定价计算器或展示型网页，但研究价值仍需由问题、证据和解释共同决定。')
    p('最可辩护的候选贡献，是一个可审计的比较协议，以及三个相互连接的条件性发现：期限匹配不自动改善净对冲；联合验证目标在小样本中没有稳定优势；有限特征CVaR学习在匹配内改善却可能在分布变化后失效。负结果能否发表，取决于其是否排除了合理替代解释、是否回答了真正未决的问题，而不是负结果本身。')
    table(['尚未建立的主张','缺少的证据或证明'],[
        ['新理论或新算法','当前使用已知Ridge、带状规则、CVaR和LP构造'],
        ['真实期权交易价值','缺少期权买卖报价、利率曲线、报价时序与可执行费用'],
        ['对强学习基线的优势','尚未实施同信息与算力预算下的deep-hedging对照'],
        ['外部确认与广泛泛化','历史设计有回顾性；2026追加期短；资产和合约有限'],
        ['全域可行性','仓位限制只在训练状态成立，实际已有样本外违反'],
        ['穷尽技术改进空间','可以继续研究，但不存在可证明的绝对技术终点']],[1.55,5.0])
    p('这些边界并不抹去工程和研究训练价值。它们决定合适的论题大小：可以讨论一个明确的有限样本方法问题，不能把一套完整软件等同于金融研究最前沿。若增加模块只为抬高代码行数，反而会稀释主要问题并扩大验证负担。')
    p('建议英文论文主线为 Maturity Alignment and Finite-Sample Selection in Cost-Constrained Hedging。正文只保留一个主要研究问题、预先指定对比和必要的失效实验；完整软件、额外费用和敏感性结果作为可复现附件，并遵守目标会场是否接受附件的规定。当前中文论文式报告是证据底稿，不是假装已发表的英文论文。')
    h('使用报告时的表述',2)
    p('可以表述为：基于金融数学课程，建立了一个在AI协助下开发的可复现期权对冲研究平台，比较期限预测、有限样本选择和尾部风险优化，分析了匹配与分布变化条件下的表现。')
    p('尚不应表述为：提出市场最优算法、完成真实交易盈利验证、独立完成全部研究、获得论文录用，或与课程讲师所在大学建立了科研合作。')

    page('16 论文和会议的可行路径')
    p('以下以2026年10月6日核验的官方页面为准。研究方向与一个渠道相关，不代表达到录用标准；摘要展示、会议论文与期刊论文必须分别记录。准备材料后仍应由实际作者审阅并决定投稿。本次没有投稿、联系主办方或支付费用。')
    h('SIAM FM27 研究展示',2)
    p('金融数学主题匹配较直接。contributed poster或lecture摘要截止2026年12月15日23:59美国东部时间，限1500字符含空格。会议于2027年6月15至18日在美国Arlington举行，要求现场参与。该摘要路径适合先获得同行反馈，不能据此写成已发表完整会议论文。')
    link(p('官方要求  '),'FM27 submissions','https://www.siam.org/conferences-events/siam-conferences/fm27/submissions/')
    h('IEEE SSCI 2027 CIFEr 后期成果展示',2)
    p('金融工程与经济计算智能方向相符。late-breaking最多2页含参考文献，poster摘要最多250词，截止2026年11月1日；页面未说明时区。完整与短论文截止已过。官方明确late-breaking和poster不进论文集，不能以IEEE正式论文表述。')
    link(p('官方要求  '),'SSCI call for papers','https://attend.ieee.org/ssci-2027/submissions/call-for-papers/')
    h('工作坊与正式论文路径',2)
    p('RAIOps4Fin 2026关注金融AI验证与治理，2026年10月12日23:59 AoE截止，最多6页，不计参考文献与附录。至少5页的录用稿可能在作者同意后被考虑纳入CEUR，并非自动收录或进入ACM主会论文集；截止很近，不宜为了赶时间省略真实作者复核。ICAIF 2026主会8月9日AoE截止已过，尚未核验2027征稿日期。')
    link(p('官方要求  '),'RAIOps4Fin workshop','https://raiops4fin2026.github.io/ICAIF/');link(p('官方要求  '),'ICAIF main track','https://icaif2026.org/call-for-papers.html')
    h('SIURO 本科研究期刊',2)
    p('若项目确为本科阶段研究，SIAM Undergraduate Research Online是值得研究的正式期刊路径。稿件须为英文PDF，通常约20页、3 MB，摘要不超过250词。真实项目指导人须出具函件，确认本科期间的研究与重要贡献，评估学生独立完成的部分，并推荐3名审稿人。指导人可以是高校教师，也可以是实际指导研究的非学术机构或政府实验室人员。软件和AI不能替代真实指导函；当前未核验具备这项条件。')
    link(p('官方要求  '),'SIURO author instructions','https://www.siam.org/publications/siam-journals/siam-undergraduate-research-online-siuro/instructions-for-authors/')

    page('17 申请能力证据与作者责任')
    p('本项目与此前十个目标硕士项目相关能力的并集相连，但不会把所有院校需求压成同一套标签。金融数学语境能展示统计、计算、优化和系统能力；对健康数据科学或生物统计，只能提供可迁移方法证据，不能替代健康问题经验、研究伦理或相应正式先修。')
    table(['能力','本项目可供讲述和核验的材料'],[
        ['概率与数学建模','P与Q的区别、自融资账本、随机波动率、CVaR凸构造'],
        ['统计推断','时间清除、配对比较、移动块敏感性、多重比较与选择不确定性'],
        ['数据科学','因果特征、训练拟合边界、真实来源、失配与泛化诊断'],
        ['优化与决策','共同预算、不可行回退、LP证书、风险与成本取舍'],
        ['计算与工程','可复现脚本、单元及产物测试、结果API、数据清单和版本隔离'],
        ['研究表达','近邻文献、预先指定比较、负结果、局限与投稿类别辨析']],[1.30,5.25])
    p('项目由杨又铭发起，与金融数学课程兴趣相连。课程资料记录Introduction to Financial Mathematics、2024年10月11日至2025年1月5日、54小时、84.40分；证书不能自动说明获得授课教师所属大学的正式学分或科研指导。项目也没有由此成为教师指导课题。')
    p('本轮研究设计、代码、实验组织与文稿整理获得实质AI协助。提交学术成果时，应按目标渠道规则具体说明使用范围，并由人类作者承担准确性和原创性责任。现阶段没有证据证明学生已独立掌握全部实现或逐项完成实验，因此不能预写“独立设计并完成全部方法”或杜撰导师、共同作者与个人工时。')
    p('适合后续记录的真实个人贡献包括：手推并解释账本和LP约束；独立复跑一个完整实验；核验一个数据清除边界；复现并说明一项负结果；读懂与论题最接近的文献；记录自己修改的方法和对应结果。这些活动完成后再写入贡献表，能够把已有软件成果转为可答辩的个人研究经历。')
    p('软件发表也有额外门槛。JOSS当前预审要求包括超过六个月的公开历史，并在该时期内持续活跃开发；这份新版本不满足立即投稿条件。会议并不天然比期刊容易，正式论文也不会因为代码多而降低原创性与证据要求。')

    page('18 结论与值得继续研究的问题')
    p('V2将课程延伸项目推进为有明确比较协议、实际结果、数学审计和发表边界的计算研究原型。它没有证实期限Ridge和联合MSE选择普遍改善对冲，也没有把CVaR学习的匹配内优势外推到变化场景。结果支持继续研究有限样本与模型风险，而不支持技术已到绝对极限或成果已获同行认可。')
    p('未来改进应围绕论文主张进行。若要强化方法评价，最有用的是在不读结果前冻结新的资产、期限和执行假设，增加完整流程重复，检查极端episode与目标尺度对选择的影响，并报告全部敏感性。若要主张新学习方法，则需要明确的新约束或结构、与强学习基线公平比较，以及可复查的理论或实验增量。')
    p('真实期权数据会改变研究问题。它要求正确处理报价时间、bid–ask、利率、股息或双币融资、到期日和合约滚动，不能简单将价格CSV换成另一份CSV。取得合适授权并定义可执行假设之前，当前实验应继续明确标为假设期权研究。')
    p('全域风险约束也是实质问题：可以研究保证仓位界限的策略参数化、约束投影与风险重新评价，或对分布变化的稳健目标。但任何新策略都要重新建立账本、优化目标和实际动作的一致性。不能一边clip动作，一边继续使用未经修改的凸最优性证书。')
    p('当前最合理的停止点是交付这套经过验证的研究版本，而不是持续扩大无关功能。它已经提供用于导师讨论、研究展示和正式稿件发展的一组具体材料。最终论文题目、贡献边界及作者贡献应在阅读近邻文献和实际复核后由研究者与真实指导者确定。')
    papers=literature['papers']
    refs=[{'id':r['id'],'citation':', '.join(r['authors'])+f' ({r["year"]}). '+r['title']+'.','url':r['url']} for r in papers]
    refs.extend([
        {'id':'R8','citation':'Rockafellar RT and Uryasev S (2000). Optimization of Conditional Value-at-Risk. Journal of Risk 2(3), 21–41.','url':'https://doi.org/10.21314/jor.2000.038'},
        {'id':'R9','citation':'Black F and Scholes M (1973). The Pricing of Options and Corporate Liabilities. Journal of Political Economy 81(3), 637–654.','url':'https://doi.org/10.1086/260062'},
        {'id':'R10','citation':'Patton AJ (2011). Volatility forecast comparison using imperfect volatility proxies. Journal of Econometrics 160(1), 246–256.','url':'https://scholars.duke.edu/publication/792433'},
        {'id':'R11','citation':'Künsch HR (1989). The Jackknife and the Bootstrap for General Stationary Observations. Annals of Statistics 17(3), 1217–1241.','url':'https://doi.org/10.1214/aos/1176347265'},
        {'id':'D1','citation':'European Central Bank. Euro foreign exchange reference rates. Archived source and modification hashes are supplied with this project.','url':'https://www.ecb.europa.eu/stats/policy_and_exchange_rates/euro_reference_exchange_rates/html/index.en.html'}])
    for offset in [0,8]:
        page('参考文献' if offset==0 else '参考文献续')
        for r in refs[offset:offset+8]:
            par=p(f'[{r["id"]}] {r["citation"]}');link(par,'  原文或出版记录',r['url'])
            if r['id'].startswith('L'):
                source=next(x for x in papers if x['id']==r['id']);p(source['relevance'])
        p('文献检索截至2026年10月6日。2026年arXiv来源明确按预印本理解，不能仅凭arXiv记录认定已同行评审。完整核验说明见docs/v2_literature.json；会议时限与规则见docs/publication_routes.json。')
    out=ROOT/'report/金融数学风险研究_V2_论文式报告与发表评估.docx';out.parent.mkdir(exist_ok=True);doc.save(out)
    evidence={'report':out.name,'input_sha256':{n:hashlib.sha256((ROOT/n).read_bytes()).hexdigest() for n in ['results/v2_historical_summary.json','results/v2_selection_summary.json','results/v2_convex_summary.json','results/validation.json','docs/v2_literature.json','docs/publication_routes.json']},'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'authoring_note':'Numbers loaded from executed evidence; conclusions are tied to the frozen V2 experiment and require review if protocol changes.'}
    (ROOT/'results/v2_report_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'path':str(out),**evidence},ensure_ascii=False))

if __name__=='__main__':main()
